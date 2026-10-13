from pathlib import Path


def make_docx(path: Path, text: str):
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    document.save(path)


def make_xlsx(path: Path, text: str):
    from xlsxwriter import Workbook

    with Workbook(path) as workbook:
        sheet = workbook.add_worksheet("Inventory")
        sheet.write(0, 0, text)


def make_pptx(path: Path, text: str):
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Product launch"
    slide.placeholders[1].text = text
    presentation.save(path)


def make_pdf(path: Path, text: str):
    text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 40 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    data = bytearray(b"%PDF-1.4\n")
    positions = [0]
    for index, obj in enumerate(objects, 1):
        positions.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(positions)}\n0000000000 65535 f \n".encode())
    for position in positions[1:]:
        data.extend(f"{position:010d} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size {len(positions)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(data)


def create_corpus(root: Path) -> dict[str, str]:
    from PIL import Image, ImageDraw

    root.mkdir(parents=True, exist_ok=True)
    nested = root / "nested"
    nested.mkdir(exist_ok=True)
    texts = {
        "orchard-notes.txt": "The apple orchard grows fresh fruit. Farmers harvest ripe apples in autumn, prune trees and manage irrigation.",
        "nested/travel.rtf": "A rail journey through the Swiss Alps. The holiday includes high mountains, alpine hiking, glaciers and snow.",
        "budget.docx": "Annual financial forecast. Revenue, expenses, quarterly budgets, cash flow projections and operating profit.",
        "launch.pptx": "Product roadmap and launch presentation. Release timeline, customer research and software delivery milestones.",
        "inventory.xlsx": "Warehouse stock inventory. Shipping counts, logistics, incoming deliveries and supply chain management.",
        "stars.pdf": "Astronomers study distant stars and galaxies. Telescopes collect light from planets and the Milky Way.",
    }
    (root / "orchard-notes.txt").write_text(texts["orchard-notes.txt"], encoding="utf-8")
    (nested / "travel.rtf").write_text(r"{\rtf1\ansi " + texts["nested/travel.rtf"] + "}", encoding="ascii")
    make_docx(root / "budget.docx", texts["budget.docx"])
    make_pptx(root / "launch.pptx", texts["launch.pptx"])
    make_xlsx(root / "inventory.xlsx", texts["inventory.xlsx"])
    make_pdf(root / "stars.pdf", texts["stars.pdf"])
    red = Image.new("RGB", (256, 256), "white")
    ImageDraw.Draw(red).ellipse((40, 40, 216, 216), fill="red")
    red.save(root / "image-a.png")
    blue = Image.new("RGB", (256, 256), "white")
    ImageDraw.Draw(blue).rectangle((40, 40, 216, 216), fill="blue")
    blue.save(root / "image-b.jpg")
    return texts
